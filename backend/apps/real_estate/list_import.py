"""LIST.xlsx / CSV の dry-run 取込（第 1 版は検証と報告だけで、取引は作らない）。

- XLSX は `工作表1` だけを読む。他のシートは開かない・列挙しない・報告にも出さない。
- 各行の元ファイル・シート・行番号・原値を保持して報告する。
- 「-」・空・0 を区別して規範化し、列ごとに件数を報告する。
- 顧客・会社・物件は候補を示すだけで、自動では結び付けない。担当者候補は既存記録の文字列だけを使う。
- 同じファイルを何度実行しても結果は同じ（書き込みは実行履歴だけ）。
"""
import csv
import hashlib
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

TARGET_SHEET = '工作表1'
MAX_ROWS = 5000

# 見出し → (項目, 型)
COLUMNS = {
    '日期': ('transaction_date', 'date'),
    '番号': ('source_number', 'text'),
    '客名': ('party_name', 'text'),
    '物件名': ('property_name', 'text'),
    '部屋番号': ('room_number', 'code'),
    '種類': ('transaction_type', 'type'),
    '管理会社': ('management_company_name', 'text'),
    '中介费': ('brokerage_fee', 'amount'),
    '广告料': ('advertising_fee', 'amount'),
    '支払い状態': ('payment_status', 'payment'),
    '支払日': ('payment_date', 'date'),
    '振込状態': ('transfer_status', 'transfer'),
    '手续费': ('handling_fee', 'amount'),
    '担当者': ('responsible_name', 'text'),
}
TYPE_VALUES = {'賃貸': 'rental', '賃貸借': 'rental', '売買': 'sale'}
PAYMENT_VALUES = {'済み': 'paid', '支払済み': 'paid', '相殺': 'offset', '未払い': 'unpaid', '未': 'unpaid'}
TRANSFER_VALUES = {'振込済み': 'transferred', '済み': 'transferred', '未': 'pending', '振込待ち': 'pending'}
DASHES = {'-', '－', 'ー', '―', '—', '–'}
# 本登録時に必要だが欠けていれば「要補充」とする（推測で埋めない）
TO_COMPLETE = (('transaction_date', '日期'), ('responsible_name', '担当者'),
               ('management_company_name', '管理会社'), ('payment_status', '支払い状態'))


class ListImportError(ValueError):
    pass


def file_sha256(content):
    return hashlib.sha256(content).hexdigest()


def _decode_csv(content):
    for encoding in ('utf-8-sig', 'cp932'):
        try:
            return content.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise ListImportError('CSV の文字コードを判定できません（UTF-8 または Shift_JIS で保存してください）。')


def read_table(filename, content):
    """(見出し, 行リスト[(行番号, 値の list)], シート名, シート一覧) を返す。"""
    name = (filename or '').lower()
    if name.endswith('.csv'):
        text, _ = _decode_csv(content)
        rows = list(csv.reader(io.StringIO(text)))
        if not rows:
            raise ListImportError('データがありません。')
        return rows[0], [(i + 2, r) for i, r in enumerate(rows[1:])], 'CSV', ['CSV']
    if name.endswith('.xlsx'):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise ListImportError('Excel ファイルを開けません。') from exc
        if TARGET_SHEET not in wb.sheetnames:
            raise ListImportError(f'シート「{TARGET_SHEET}」がありません（このシートだけを読み込みます）。')
        ws = wb[TARGET_SHEET]
        iterator = ws.iter_rows(values_only=True)
        headers = next(iterator, None)
        if headers is None:
            raise ListImportError('データがありません。')
        rows = []
        for index, values in enumerate(iterator, start=2):
            rows.append((index, list(values)))
            if len(rows) > MAX_ROWS:
                raise ListImportError(f'{MAX_ROWS} 行を超えるファイルは読み込めません。')
        return [str(h).strip() if h is not None else '' for h in headers], rows, TARGET_SHEET, [TARGET_SHEET]
    raise ListImportError('CSV または XLSX ファイルを選択してください。')


def raw_text(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d') if value.time() == datetime.min.time() else value.isoformat(sep=' ')
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def classify(value):
    """'empty' / 'dash' / 'zero' / 'value'。"""
    text = raw_text(value)
    if text == '':
        return 'empty'
    if text in DASHES:
        return 'dash'
    try:
        if Decimal(text.replace(',', '')) == 0:
            return 'zero'
    except InvalidOperation:
        pass
    return 'value'


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = raw_text(value)
    for pattern in ('%Y-%m-%d', '%Y/%m/%d', '%Y.%m.%d', '%Y年%m月%d日'):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    match = re.fullmatch(r'(\d{1,2})[/.月](\d{1,2})日?', text)
    if match:
        return 'no_year'
    return None


def parse_amount(value):
    text = raw_text(value).replace(',', '').replace('円', '').replace('￥', '').replace('¥', '')
    try:
        amount = Decimal(text)
    except InvalidOperation:
        return None
    if amount < 0 or amount != amount.to_integral_value():
        return None
    return int(amount)


def normalize_row(headers, values):
    raw, normalized, errors, notes = {}, {}, [], []
    for index, header in enumerate(headers):
        spec = COLUMNS.get(header)
        value = values[index] if index < len(values) else None
        if spec is None:
            continue
        field, kind = spec
        raw[header] = raw_text(value)
        state = classify(value)
        if state in ('empty', 'dash'):
            normalized[field] = None
            if state == 'dash':
                notes.append({'column': header, 'code': 'dash', 'message': '「-」を空として扱いました（元の値は保持）'})
            continue
        if kind == 'date':
            parsed = parse_date(value)
            if parsed == 'no_year':
                errors.append({'column': header, 'raw': raw[header], 'message': '年が無い日付は判定しません（要確認）'})
                normalized[field] = None
            elif parsed is None:
                errors.append({'column': header, 'raw': raw[header], 'message': '日付として読めません'})
                normalized[field] = None
            else:
                normalized[field] = parsed.isoformat()
        elif kind == 'amount':
            amount = parse_amount(value)
            if amount is None:
                errors.append({'column': header, 'raw': raw[header], 'message': '金額として読めません（0 以上の整数）'})
            normalized[field] = amount
            if state == 'zero':
                notes.append({'column': header, 'code': 'zero', 'message': '0 円として記録されています（空とは区別）'})
        elif kind == 'type':
            normalized[field] = TYPE_VALUES.get(raw[header])
            if normalized[field] is None:
                errors.append({'column': header, 'raw': raw[header], 'message': '種類は「賃貸」「売買」のいずれかです'})
        elif kind == 'payment':
            normalized[field] = PAYMENT_VALUES.get(raw[header])
            if normalized[field] is None:
                errors.append({'column': header, 'raw': raw[header], 'message': '支払い状態の値が不明です'})
        elif kind == 'transfer':
            normalized[field] = TRANSFER_VALUES.get(raw[header])
            if normalized[field] is None:
                errors.append({'column': header, 'raw': raw[header], 'message': '振込状態の値が不明です'})
        elif kind == 'code':
            normalized[field] = raw[header]
        else:
            normalized[field] = raw[header]
    return raw, normalized, errors, notes


def column_stats(headers, rows):
    stats = {}
    for index, header in enumerate(headers):
        if header not in COLUMNS:
            continue
        counter = {'empty': 0, 'dash': 0, 'zero': 0, 'value': 0}
        for _, values in rows:
            counter[classify(values[index] if index < len(values) else None)] += 1
        stats[header] = counter
    return stats


def _key(*parts):
    return '|'.join(re.sub(r'\s+', '', str(p or '')).lower() for p in parts)


def build_report(filename, content, *, candidates):
    """dry-run の報告を作る。candidates は候補検索の関数群（範囲は呼び出し側の権限で絞る）。"""
    headers, rows, sheet, sheets = read_table(filename, content)
    missing_columns = [h for h in COLUMNS if h not in headers]
    data_rows = [(n, v) for n, v in rows if any(classify(x) != 'empty' for x in v)]
    blank_rows = len(rows) - len(data_rows)
    results = []
    seen = {}
    for row_number, values in data_rows:
        raw, norm, errors, notes = normalize_row(headers, values)
        to_complete = [label for field, label in TO_COMPLETE if not norm.get(field)]
        if not norm.get('party_name'):
            errors.append({'column': '客名', 'raw': raw.get('客名', ''), 'message': '客名がありません'})
        if not norm.get('property_name'):
            errors.append({'column': '物件名', 'raw': raw.get('物件名', ''), 'message': '物件名がありません'})
        if not norm.get('transaction_type') and not any(e['column'] == '種類' for e in errors):
            notes.append({'column': '種類', 'code': 'default_type', 'message': '種類が空のため、本登録時は賃貸として扱う予定（要確認）'})
        key = _key(norm.get('source_number')) if norm.get('source_number') else _key(
            norm.get('party_name'), norm.get('property_name'), norm.get('room_number'), norm.get('transaction_date'))
        duplicates_in_file = list(seen.get(key, []))
        seen.setdefault(key, []).append(row_number)
        results.append({
            'row_number': row_number,
            'source': {'file': filename, 'sheet': sheet, 'row': row_number},
            'raw': raw,
            'values': norm,
            'errors': errors,
            'notes': notes,
            'to_complete': to_complete,
            'duplicate_in_file_rows': list(duplicates_in_file),
            'duplicate_existing': candidates['existing'](norm),
            'candidates': {
                'customer': candidates['customer'](norm.get('party_name')),
                'management_company': candidates['company'](norm.get('management_company_name')),
                'property': candidates['property'](norm.get('property_name'), norm.get('room_number')),
                'responsible': candidates['responsible'](norm.get('responsible_name')),
            },
        })
    summary = {
        'rows': len(data_rows),
        'blank_rows_skipped': blank_rows,
        'error_rows': sum(1 for r in results if r['errors']),
        'to_complete_rows': sum(1 for r in results if r['to_complete']),
        'duplicate_in_file_rows': sum(1 for r in results if r['duplicate_in_file_rows']),
        'duplicate_existing_rows': sum(1 for r in results if r['duplicate_existing']),
        'ready_rows': sum(1 for r in results if not r['errors'] and not r['to_complete']
                          and not r['duplicate_in_file_rows'] and not r['duplicate_existing']),
    }
    return {
        'file_name': filename,
        'sheet': sheet,
        'sheets_found': sheets,
        'headers': [header for header in headers if header in COLUMNS],
        'missing_columns': missing_columns,
        'column_stats': column_stats(headers, data_rows),
        'summary': summary,
        'results': results,
    }


def error_report_csv(report):
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(['行番号', '種別', '列', '元の値', '内容'])
    for r in report.get('results', []):
        for e in r['errors']:
            writer.writerow([r['row_number'], '誤り', e['column'], e.get('raw', ''), e['message']])
        for label in r['to_complete']:
            writer.writerow([r['row_number'], '要補充', label, r['raw'].get(label, ''), '本登録前に補充が必要'])
        if r['duplicate_in_file_rows']:
            writer.writerow([r['row_number'], '重複（ファイル内）', '', '', '同じ内容の行：' + '、'.join(map(str, r['duplicate_in_file_rows']))])
        if r['duplicate_existing']:
            writer.writerow([r['row_number'], '重複（登録済み）', '', '', '既存：' + '、'.join(d['number'] for d in r['duplicate_existing'])])
        for n in r['notes']:
            writer.writerow([r['row_number'], '注記', n['column'], r['raw'].get(n['column'], ''), n['message']])
    return ('﻿' + buffer.getvalue()).encode('utf-8')
