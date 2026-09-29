"""P3 LIST.xlsx dry-run：工作表1 だけ・强哥は読まない、-/空/0 の区別、3 つの金額は別々、候補のみ、重複、再実行。"""
import io
import json
from datetime import datetime

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from openpyxl import Workbook

from apps.audit.models import AuditLog
from apps.companies.models import Company
from apps.customers.models import Customer
from apps.real_estate.models import RealEstateTransaction
from apps.real_estate.tests import RealEstateFixture

HEADERS = ['日期', '番号', '客名', '物件名', '部屋番号', '種類', '管理会社', '中介费', '广告料', '支払い状態', '支払日',
           '向SUNRISE請求書金額', '向客人請求金額', 'SUNRISE請求書金額', '振込状態', '手续费', '担当者']
SENTINEL = 'QIANGGE-SHOULD-NEVER-APPEAR'


def workbook_bytes(rows, extra_sheets=True):
    wb = Workbook()
    ws = wb.active
    ws.title = '工作表1'
    ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    if extra_sheets:
        wb.create_sheet('强哥').append([SENTINEL, 999999])
        wb.create_sheet('工作表2').append(['メモ'])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


ROWS = [
    [datetime(2026, 4, 3), 'A-001', '王 一郎', 'サンライズ天王寺', 301, '賃貸', '大阪管理', 85000, 42500, '済み',
     datetime(2026, 4, 20), 127500, 93500, 110000, '振込済み', None, '担当A'],
    [None, 'A-002', '張 二郎', 'ハイツ勝山', 102, None, None, '-', 0, None, None, None, None, None, None, None, None],
    ['4/5', 'A-003', '李 三', 'ハイツ勝山', '102号', '賃貸', '京都管理', 'abc', 10000, '相殺', None, None, None, None,
     None, 500, '担当A'],
    [None] * 17,
    [datetime(2026, 4, 3), 'A-001', '王 一郎', 'サンライズ天王寺', 301, '賃貸', '大阪管理', 85000, 42500, '済み', None,
     None, None, None, None, None, '担当A'],
    [datetime(2026, 5, 1), 'A-005', None, None, None, '売買', '大阪管理', 1, 1, '済み', None, None, None, None, None, None, '担当A'],
]


@override_settings(REAL_ESTATE_IMPORT_DRY_RUN_ENABLED=True)
class ListDryRunTests(RealEstateFixture, TestCase):
    def dry_run(self, user, content=None, name='LIST.xlsx'):
        self.client.force_login(user)
        return self.client.post('/api/real-estate/imports/dry-run/', {
            'file': SimpleUploadedFile(name, content if content is not None else workbook_bytes(ROWS))})

    def test_reads_only_sheet1_and_never_qiangge(self):
        response = self.dry_run(self.li)
        self.assertEqual(response.status_code, 201, response.content)
        body = response.json()
        text = json.dumps(body, ensure_ascii=False)
        self.assertEqual(body['sheet'], '工作表1')
        self.assertNotIn(SENTINEL, text)
        self.assertNotIn('强哥', text)
        self.assertEqual(body['summary']['rows'], 5)
        self.assertEqual(body['summary']['blank_rows_skipped'], 1)
        self.assertTrue(body['dry_run'])
        self.assertEqual(RealEstateTransaction.objects.count(), 0)

    def test_normalization_amounts_candidates_duplicates(self):
        Customer.objects.create(name='王 一郎', birth_date='1990-01-01')
        Company.objects.create(name='大阪管理株式会社')
        body = self.dry_run(self.li).json()
        rows = {r['row_number']: r for r in body['results']}
        first = rows[2]
        self.assertEqual(first['source'], {'file': 'LIST.xlsx', 'sheet': '工作表1', 'row': 2})
        self.assertEqual(first['values']['transaction_date'], '2026-04-03')
        self.assertEqual(first['values']['room_number'], '301')
        # 近い 3 つの金額は合算せず別々
        self.assertEqual(first['values']['source_billed_to_sunrise_amount'], 127500)
        self.assertEqual(first['values']['source_billed_to_client_amount'], 93500)
        self.assertEqual(first['values']['source_sunrise_invoice_amount'], 110000)
        self.assertEqual(first['values']['payment_status'], 'paid')
        self.assertEqual(first['values']['transfer_status'], 'transferred')
        self.assertEqual(first['raw']['部屋番号'], '301')
        self.assertEqual([c['name'] for c in first['candidates']['customer']], ['王 一郎'])
        self.assertEqual([c['name'] for c in first['candidates']['management_company']], ['大阪管理株式会社'])
        self.assertEqual(first['candidates']['responsible'][0]['name'], 'A')
        # -・0・空の区別
        second = rows[3]
        self.assertIsNone(second['values']['brokerage_fee'])
        self.assertEqual(second['values']['advertising_fee'], 0)
        self.assertEqual({n['code'] for n in second['notes']} >= {'dash', 'zero', 'default_type'}, True)
        self.assertEqual(set(second['to_complete']), {'日期', '担当者', '管理会社', '支払い状態'})
        stats = body['column_stats']
        self.assertEqual(stats['中介费']['dash'], 1)
        self.assertEqual(stats['广告料']['zero'], 1)
        self.assertEqual(stats['手续费']['empty'], 4)
        # 年の無い日付・読めない金額は誤り（推測しない）
        third = rows[4]
        self.assertEqual({e['column'] for e in third['errors']}, {'日期', '中介费'})
        # ファイル内の重複（同じ番号）
        self.assertEqual(rows[6]['duplicate_in_file_rows'], [2])
        # 客名・物件名が無い行
        self.assertEqual({e['column'] for e in rows[7]['errors']}, {'客名', '物件名'})
        self.assertEqual(rows[7]['values']['transaction_type'], 'sale')

    def test_repeatable_existing_duplicates_and_error_report(self):
        first = self.dry_run(self.li).json()
        self.create_tx(user=self.li, party_name='王 一郎', property_name='サンライズ天王寺', room_number='301')
        second = self.dry_run(self.li).json()
        self.assertEqual(second['previous_runs'], [first['id']])
        self.assertEqual(second['results'][0]['duplicate_existing'][0]['number'][:3], 'RE-')
        self.client.force_login(self.li)
        report = self.client.get(f"/api/real-estate/imports/{second['id']}/error-report/")
        self.assertEqual(report.status_code, 200)
        text = report.content.decode('utf-8-sig')
        self.assertIn('重複（登録済み）', text)
        self.assertIn('年が無い日付', text)
        self.assertEqual(AuditLog.objects.filter(action='import_dry_run').count(), 2)
        runs = self.client.get('/api/real-estate/imports/').json()
        self.assertEqual(len(runs), 2)

    def test_csv_and_missing_sheet(self):
        csv = ('﻿' + ','.join(HEADERS) + '\n2026-04-03,C-1,客,物件,1,賃貸,管理,-,0,済み,,,,,,,担当A\n').encode('utf-8')
        body = self.dry_run(self.li, csv, 'list.csv').json()
        self.assertEqual(body['sheet'], 'CSV')
        self.assertEqual(body['summary']['rows'], 1)
        wb = Workbook()
        wb.active.title = '别的'
        buf = io.BytesIO()
        wb.save(buf)
        self.assertEqual(self.dry_run(self.li, buf.getvalue()).status_code, 400)

    def test_permissions_and_disabled_environment(self):
        self.assertEqual(self.dry_run(self.staff_a).status_code, 403)
        self.assertEqual(self.dry_run(self.jiao).status_code, 403)
        with override_settings(REAL_ESTATE_IMPORT_DRY_RUN_ENABLED=False):
            self.assertEqual(self.dry_run(self.li).status_code, 403)
