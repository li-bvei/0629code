"""返签 visa 表の CSV/XLSX 一括取込：読込・列の対応付け・変換・行ごとの検証・重複検出。

- 読込：CSV は utf-8(BOM) → cp932 → gb18030 の順で判定（指定も可）。XLSX は複数シート対応。
- 先頭の空でない行を見出し行とし、空行は数えて飛ばす。
- 値は「元の値（raw）」と「変換後の値（value）」を分けて返し、誤りは黙って飛ばさない。
- 行データは DB に保存しない（誤り行だけ、修正用に項目・元の値・理由を batch.results に残す）。
"""
import csv
import hashlib
import io
import re
import unicodedata
from datetime import date, datetime, timedelta

from django.core.validators import validate_email
from django.core.exceptions import ValidationError

from .models import VisaReturnApplication

MAX_ROWS = 500
MAX_FILE_BYTES = 5 * 1024 * 1024
CSV_ENCODINGS = ('utf-8-sig', 'cp932', 'gb18030')

# 取込先の項目：(キー, 表示名, 種別, 必須, 見出しの別名)
FIELDS = [
    ('applicant_name', '申請人氏名', 'text', True, ['氏名', '姓名', '申请人姓名', '申請人氏名', '名前', 'name', 'applicantname']),
    ('nationality', '国籍', 'text', False, ['国籍', 'nationality']),
    ('birth_date', '生年月日', 'date', False, ['生年月日', '出生日期', '生日', 'birthdate', 'birthday', 'dateofbirth']),
    ('gender', '性別', 'gender', False, ['性別', '性别', 'gender', 'sex']),
    ('marital_status', '婚姻状況', 'marital', False, ['婚姻状況', '婚姻状况', '婚姻', 'maritalstatus']),
    ('passport_number', '旅券番号', 'code', True, ['旅券番号', 'パスポート番号', '护照号', '护照号码', 'passport', 'passportnumber', 'passportno']),
    ('passport_issue_date', '旅券発行日', 'date', False, ['旅券発行日', '护照签发日期', '签发日期', 'passportissuedate', 'issuedate']),
    ('passport_expiry_date', '旅券期限', 'date', False, ['旅券期限', '旅券有効期限', '护照有效期', '有效期', 'passportexpirydate', 'expirydate']),
    ('residence_status', '在留資格', 'text', False, ['在留資格', '在留资格', 'residencestatus']),
    ('address', '住所', 'text', False, ['住所', '地址', 'address']),
    ('phone', '電話番号', 'code', False, ['電話番号', '電話', '电话', '手机', '手机号', 'phone', 'tel']),
    ('email', 'メール', 'email', False, ['メール', 'メールアドレス', '邮箱', '电子邮件', 'email', 'mail']),
    ('occupation', '職業', 'text', False, ['職業', '职业', 'occupation', 'job']),
    ('guarantor_name', '保証人氏名', 'text', False, ['保証人氏名', '保証人', '担保人', '担保人姓名', 'guarantorname', 'guarantor']),
    ('guarantor_phone', '保証人電話番号', 'code', False, ['保証人電話番号', '担保人电话', 'guarantorphone']),
    ('guarantor_address', '保証人住所', 'text', False, ['保証人住所', '担保人地址', 'guarantoraddress']),
    ('guarantor_relationship', '申請人との関係', 'text', False, ['申請人との関係', '关系', '与申请人关系', 'relationship']),
    ('guarantor_occupation', '保証人職業', 'text', False, ['保証人職業', '担保人职业', 'guarantoroccupation']),
    ('note', '備考', 'text', False, ['備考', '备注', 'note', 'remarks']),
    ('form_data.pinyin_name1', '拼音（姓）', 'text', False, ['拼音姓', 'ピンイン姓', 'pinyinname1', 'surname']),
    ('form_data.pinyin_name2', '拼音（名）', 'text', False, ['拼音名', 'ピンイン名', 'pinyinname2', 'givenname']),
    ('form_data.birth_place', '出生地', 'text', False, ['出生地', '出生地点', 'birthplace']),
    ('form_data.chinese_id', '身分証番号', 'code', False, ['身份证号', '身份证号码', '身分証番号', 'chineseid', 'idnumber']),
    ('form_data.entry_port', '入国予定港', 'text', False, ['入国港', '入境口岸', 'entryport']),
    ('form_data.airline', '航空会社・便名', 'text', False, ['航空会社', '航班', 'airline', 'flight']),
    ('form_data.entry_time1', '入国予定日', 'date', False, ['入国予定日', '入境日期', 'entrydate', 'entrytime1']),
    ('form_data.current_address', '現住所（中国）', 'text', False, ['现住址', '現住所', 'currentaddress']),
    ('form_data.workplace_name', '勤務先', 'text', False, ['勤務先', '工作单位', 'workplace', 'workplacename']),
    ('form_data.workplace_phone', '勤務先電話', 'code', False, ['勤務先電話', '单位电话', 'workplacephone']),
]
FIELD_KEYS = [f[0] for f in FIELDS]
FIELD_META = {f[0]: {'label': f[1], 'type': f[2], 'required': f[3]} for f in FIELDS}
GENDER_MAP = {'male': 'male', 'm': 'male', '男': 'male', '男性': 'male',
              'female': 'female', 'f': 'female', '女': 'female', '女性': 'female'}
MARITAL_MAP = {'single': 'single', '未婚': 'single', 'married': 'married', '既婚': 'married', '已婚': 'married',
               'divorced': 'divorced', '離婚': 'divorced', '离婚': 'divorced', '离异': 'divorced',
               'widowed': 'widowed', '死別': 'widowed', '丧偶': 'widowed'}


class VisaImportError(Exception):
    pass


def _key(value):
    text = unicodedata.normalize('NFKC', str(value or '')).strip().lower()
    return re.sub(r'[\s　_\-・（）()［］\[\]:：]', '', text)


ALIAS_INDEX = {}
for _field, _label, _type, _required, _aliases in FIELDS:
    for _alias in [_label, _field.split('.')[-1], *_aliases]:
        ALIAS_INDEX.setdefault(_key(_alias), _field)


def file_sha256(content):
    return hashlib.sha256(content).hexdigest()


# --- 読込 ------------------------------------------------------------------------

def _cell_to_raw(value):
    """セル値を (表示用の文字列, 元の型情報) にする。"""
    if value is None:
        return '', None
    if isinstance(value, datetime):
        return value.date().isoformat(), 'date'
    if isinstance(value, date):
        return value.isoformat(), 'date'
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE', 'bool'
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value)), 'number'
        return str(value), 'number'
    if isinstance(value, int):
        return str(value), 'number'
    return str(value).strip(), 'text'


def _rows_to_table(rows):
    """先頭の空でない行を見出しに、以降の空行を飛ばして表にする。"""
    headers, body, blank_rows, header_row_number = None, [], 0, None
    for index, row in enumerate(rows, start=1):
        cells = [_cell_to_raw(v) for v in row]
        if not any(text for text, _ in cells):
            if headers is not None:
                blank_rows += 1
            continue
        if headers is None:
            headers = [text or f'列{i + 1}' for i, (text, _) in enumerate(cells)]
            header_row_number = index
            continue
        if len(body) >= MAX_ROWS:
            raise VisaImportError(f'一度に取り込めるのは {MAX_ROWS} 行までです。ファイルを分割してください。')
        values = {}
        types = {}
        for i, header in enumerate(headers):
            text, kind = cells[i] if i < len(cells) else ('', None)
            values[header] = text
            if kind:
                types[header] = kind
        body.append({'row_number': index, 'cells': values, 'cell_types': types})
    if headers is None:
        raise VisaImportError('見出し行が見つかりません（空のファイルです）。')
    return {'headers': headers, 'rows': body, 'blank_rows_skipped': blank_rows, 'header_row_number': header_row_number}


def read_csv(content, encoding=None):
    candidates = [encoding] if encoding else CSV_ENCODINGS
    last_error = None
    for name in candidates:
        try:
            text = content.decode(name)
        except (UnicodeDecodeError, LookupError) as exc:
            last_error = exc
            continue
        dialect = csv.excel
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=',\t;')
        except csv.Error:
            pass
        table = _rows_to_table(csv.reader(io.StringIO(text), dialect))
        return {**table, 'encoding': name, 'sheets': [], 'sheet_name': ''}
    raise VisaImportError(f'文字コードを判定できません（UTF-8 / Shift_JIS / GB18030 に対応）。{last_error or ""}'.strip())


def read_xlsx(content, sheet_name=None):
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # 破損・形式違い
        raise VisaImportError('Excel ファイルを読み込めません（.xlsx 形式か確認してください）。') from exc
    sheets = workbook.sheetnames
    target = sheet_name or sheets[0]
    if target not in sheets:
        raise VisaImportError(f'シート「{target}」がありません。')
    table = _rows_to_table(workbook[target].iter_rows(values_only=True))
    workbook.close()
    return {**table, 'encoding': '', 'sheets': sheets, 'sheet_name': target}


def read_upload(file_name, content, *, sheet_name=None, encoding=None):
    if len(content) > MAX_FILE_BYTES:
        raise VisaImportError('ファイルが大きすぎます（5MB まで）。')
    lower = (file_name or '').lower()
    if lower.endswith('.xlsx'):
        return read_xlsx(content, sheet_name)
    if lower.endswith('.csv') or lower.endswith('.txt') or lower.endswith('.tsv'):
        return read_csv(content, encoding)
    raise VisaImportError('CSV または XLSX ファイルを選択してください。')


def auto_mapping(headers):
    """見出し → 取込先項目の自動対応付け（同じ項目への重複は最初の列だけ）。"""
    mapping, used = {}, set()
    for header in headers:
        field = ALIAS_INDEX.get(_key(header))
        if field and field not in used:
            mapping[header] = field
            used.add(field)
    return mapping


# --- 変換・検証 --------------------------------------------------------------------

EXCEL_EPOCH = date(1899, 12, 30)


def parse_date_value(raw, kind=None):
    text = unicodedata.normalize('NFKC', str(raw or '')).strip()
    if not text:
        return None
    if kind == 'number' and re.fullmatch(r'\d{5}', text):
        return EXCEL_EPOCH + timedelta(days=int(text))  # Excel のシリアル値
    match = re.fullmatch(r'(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})日?', text)
    if match:
        y, m, d = map(int, match.groups())
    elif re.fullmatch(r'\d{8}', text):
        y, m, d = int(text[:4]), int(text[4:6]), int(text[6:])
    else:
        raise ValueError
    return date(y, m, d)


def convert_value(field, raw, kind=None):
    """(値, 誤り, 警告) を返す。"""
    meta = FIELD_META[field]
    text = unicodedata.normalize('NFKC', str(raw or '')).strip()
    field_type = meta['type']
    if not text:
        if meta['required']:
            return None, f'{meta["label"]}は必須です。', None
        return ('' if field_type not in ('date',) else None), None, None
    if field_type == 'date':
        try:
            return parse_date_value(text, kind).isoformat(), None, None
        except (ValueError, TypeError):
            return None, f'{meta["label"]}の日付形式が正しくありません（例：2026-09-01）。', None
    if field_type == 'gender':
        value = GENDER_MAP.get(text.lower())
        return (value, None, None) if value else (None, f'{meta["label"]}は「男性/女性」で入力してください。', None)
    if field_type == 'marital':
        value = MARITAL_MAP.get(text.lower())
        return (value, None, None) if value else (None, f'{meta["label"]}は「未婚/既婚/離婚/死別」で入力してください。', None)
    if field_type == 'email':
        try:
            validate_email(text)
        except ValidationError:
            return None, f'{meta["label"]}の形式が正しくありません。', None
        return text, None, None
    if field_type == 'code':
        # 旅券番号・電話などは文字列のまま扱う（先頭の 0 を保持）。
        value = re.sub(r'\s+', '', text) if field.endswith('passport_number') else text
        if field.endswith('passport_number'):
            value = value.upper()
        warning = None
        if kind == 'number':
            warning = f'{meta["label"]}が数値として保存されていたため、先頭の 0 が失われている可能性があります。元の値を確認してください。'
        return value, None, warning
    return text, None, None


def _existing_passports(passports):
    rows = VisaReturnApplication.objects.filter(passport_number__in=passports).values_list('passport_number', 'applicant_name', 'id')
    found = {}
    for passport, name, app_id in rows:
        found.setdefault(passport.upper(), []).append({'applicant_name': name, 'id': app_id})
    return found


def validate_rows(rows, mapping, shared=None):
    """行ごとに変換・検証し、重複（既存の申請・ファイル内）を警告として付ける。"""
    shared = {k: v for k, v in (shared or {}).items() if k in FIELD_META}
    unknown = sorted({field for field in mapping.values() if field and field not in FIELD_META})
    if unknown:
        raise VisaImportError(f'未知の取込先項目があります：{", ".join(unknown)}')
    missing_required = [FIELD_META[f]['label'] for f in FIELD_KEYS
                        if FIELD_META[f]['required'] and f not in mapping.values() and not shared.get(f)]
    results, seen = [], {}
    for row in rows:
        cells = row.get('cells') or {}
        kinds = row.get('cell_types') or {}
        values, errors, warnings, raw_by_field = {}, [], [], {}
        for header, field in mapping.items():
            if not field:
                continue
            raw_by_field[field] = cells.get(header, '')
            value, error, warning = convert_value(field, cells.get(header, ''), kinds.get(header))
            if error:
                errors.append({'field': field, 'label': FIELD_META[field]['label'], 'raw': cells.get(header, ''), 'message': error})
            if warning:
                warnings.append({'field': field, 'code': 'leading_zero', 'message': warning})
            values[field] = value
        for field, shared_raw in shared.items():
            if values.get(field) in (None, ''):
                value, error, _ = convert_value(field, shared_raw)
                if error:
                    errors.append({'field': field, 'label': FIELD_META[field]['label'], 'raw': shared_raw,
                                   'message': f'（共通項目）{error}'})
                values[field] = value
                raw_by_field.setdefault(field, shared_raw)
        for label in missing_required:
            errors.append({'field': '', 'label': label, 'raw': '', 'message': f'{label}の列が対応付けられていません。'})
        passport = (values.get('passport_number') or '').upper()
        if passport:
            if passport in seen:
                warnings.append({'field': 'passport_number', 'code': 'duplicate_in_file',
                                 'message': f'同じ旅券番号が {seen[passport]} 行目にもあります。'})
            else:
                seen[passport] = row.get('row_number')
        results.append({'row_number': row.get('row_number'), 'values': values, 'raw': raw_by_field,
                        'errors': errors, 'warnings': warnings})
    existing = _existing_passports([p for p in seen])
    for result in results:
        passport = (result['values'].get('passport_number') or '').upper()
        if passport in existing:
            names = '、'.join(sorted({e['applicant_name'] for e in existing[passport] if e['applicant_name']})) or '（氏名なし）'
            result['warnings'].append({'field': 'passport_number', 'code': 'duplicate_existing',
                                       'message': f'同じ旅券番号の申請が既に登録されています（{names}）。',
                                       'existing_ids': [e['id'] for e in existing[passport]]})
    return results


def to_application_payload(values):
    payload, form_data = {}, {}
    for field, value in values.items():
        if field.startswith('form_data.'):
            if value not in (None, ''):
                form_data[field.split('.', 1)[1]] = value
        elif value is not None:
            payload[field] = value
    payload['form_data'] = form_data
    return payload


def template_csv():
    headers = [FIELD_META[f]['label'] for f in FIELD_KEYS]
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(headers)
    return ('﻿' + buffer.getvalue()).encode('utf-8')


def error_report_csv(results):
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(['行番号', '項目', '元の値', '誤り'])
    for row_number, result in sorted(results.items(), key=lambda item: int(item[0])):
        for error in result.get('errors', []):
            writer.writerow([row_number, error.get('label', ''), error.get('raw', ''), error.get('message', '')])
    return ('﻿' + buffer.getvalue()).encode('utf-8')
