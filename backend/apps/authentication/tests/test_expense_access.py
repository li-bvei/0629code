"""Expense（報銷）のデータ分離：一覧・詳細・変更・作成・Excel・summary・会計ダッシュボード・
図表・copy-expenses を、未ログイン／本人／他人／業務管理者（焦・周相当）／superuser のみ／
李相当で検証する。"""
from datetime import date
from decimal import Decimal

from django.test import TestCase

from apps.accounting.models import AccountingProject, Expense, IncomeSource
from apps.audit.models import AuditLog
from apps.authentication.roles import (
    ACCOUNTING_ADMIN,
    BUSINESS_ADMIN,
    EXPENSE_VIEWER,
    STAFF,
    SYSTEM_ADMIN,
)
from apps.authentication.testing import make_user


class ExpenseAccessTests(TestCase):
    def setUp(self):
        self.li = make_user('li', roles=[SYSTEM_ADMIN, ACCOUNTING_ADMIN, BUSINESS_ADMIN], employee_name='李', superuser=True)
        self.jiao = make_user('jiao_like', roles=[BUSINESS_ADMIN, EXPENSE_VIEWER], employee_name='焦', superuser=True)
        self.staff_a = make_user('staff_a', roles=[STAFF], employee_name='A')
        self.staff_b = make_user('staff_b', roles=[STAFF], employee_name='B')
        self.su_only = make_user('su_only', superuser=True)
        self.exp_a = self._expense(self.staff_a, '1000', 'A交通費')
        self.exp_b = self._expense(self.staff_b, '2000', 'B交通費')
        self.exp_li = self._expense(self.li, '3000', '李交通費')
        self.exp_legacy = Expense.objects.create(
            expense_date=date(2026, 7, 5), category='旧データ', amount=Decimal('4000'),
        )
        IncomeSource.objects.create(source_date=date(2026, 7, 1), source_target='入金', amount=Decimal('10000'))

    def _expense(self, owner, amount, category):
        return Expense.objects.create(
            expense_date=date(2026, 7, 2), category=category, amount=Decimal(amount), owner=owner, created_by=owner,
        )

    def _ids(self, response):
        data = response.json()
        rows = data['results'] if isinstance(data, dict) and 'results' in data else data
        return {row['id'] for row in rows}

    def get(self, user, url):
        self.client.logout()
        if user is not None:
            self.client.force_login(user)
        return self.client.get(url)

    # --- 一覧・詳細 -----------------------------------------------------------
    def test_anonymous_is_rejected(self):
        self.assertIn(self.get(None, '/api/accounting/expenses/').status_code, (401, 403))

    def test_staff_sees_only_own(self):
        self.assertEqual(self._ids(self.get(self.staff_a, '/api/accounting/expenses/')), {self.exp_a.id})
        self.assertEqual(self._ids(self.get(self.staff_b, '/api/accounting/expenses/')), {self.exp_b.id})

    def test_superuser_without_business_group_cannot_read_expenses(self):
        response = self.get(self.su_only, '/api/accounting/expenses/')
        self.assertEqual(response.status_code, 403)
        for url in (
            f'/api/accounting/expenses/{self.exp_a.id}/',
            '/api/accounting/expenses/summary/',
            '/api/accounting/expenses/excel/',
            '/api/accounting/dashboard/',
            '/api/accounting/income-sources/',
        ):
            self.assertIn(self.get(self.su_only, url).status_code, (403, 404), url)

    def test_business_admin_views_all_read_only(self):
        ids = self._ids(self.get(self.jiao, '/api/accounting/expenses/'))
        self.assertEqual(ids, {self.exp_a.id, self.exp_b.id, self.exp_li.id, self.exp_legacy.id})
        self.assertTrue(AuditLog.objects.filter(action='expense_view_all', user=self.jiao).exists())
        self.assertEqual(self.get(self.jiao, f'/api/accounting/expenses/{self.exp_a.id}/').status_code, 200)

    def test_retrieve_other_is_404_for_staff(self):
        self.assertEqual(self.get(self.staff_a, f'/api/accounting/expenses/{self.exp_b.id}/').status_code, 404)
        self.assertEqual(self.get(self.staff_a, f'/api/accounting/expenses/{self.exp_legacy.id}/').status_code, 404)

    # --- 変更 ---------------------------------------------------------------
    def test_staff_cannot_change_or_delete_other(self):
        self.client.force_login(self.staff_a)
        url = f'/api/accounting/expenses/{self.exp_b.id}/'
        self.assertEqual(self.client.patch(url, {'note': 'x'}, content_type='application/json').status_code, 404)
        self.assertEqual(self.client.delete(url).status_code, 404)
        self.exp_b.refresh_from_db()
        self.assertEqual(self.exp_b.note, '')

    def test_business_admin_cannot_change_other_but_can_change_own(self):
        own = self._expense(self.jiao, '500', '焦')
        self.client.force_login(self.jiao)
        response = self.client.patch(f'/api/accounting/expenses/{self.exp_a.id}/', {'note': 'x'}, content_type='application/json')
        self.assertEqual(response.status_code, 403)
        self.assertEqual(self.client.delete(f'/api/accounting/expenses/{self.exp_a.id}/').status_code, 403)
        self.assertTrue(AuditLog.objects.filter(action='access_denied', user=self.jiao, result='denied').exists())
        response = self.client.patch(f'/api/accounting/expenses/{own.id}/', {'note': 'ok'}, content_type='application/json')
        self.assertEqual(response.status_code, 200)
        own.refresh_from_db()
        self.assertEqual(own.updated_by, self.jiao)

    def test_business_admin_cannot_change_legacy_unowned(self):
        self.client.force_login(self.jiao)
        response = self.client.patch(f'/api/accounting/expenses/{self.exp_legacy.id}/', {'note': 'x'}, content_type='application/json')
        self.assertEqual(response.status_code, 403)

    def test_li_can_change_other_with_audit(self):
        self.client.force_login(self.li)
        response = self.client.patch(f'/api/accounting/expenses/{self.exp_a.id}/', {'note': '修正'}, content_type='application/json')
        self.assertEqual(response.status_code, 200)
        self.exp_a.refresh_from_db()
        self.assertEqual(self.exp_a.note, '修正')
        self.assertEqual(self.exp_a.owner, self.staff_a)
        self.assertEqual(self.exp_a.updated_by, self.li)
        self.assertTrue(AuditLog.objects.filter(action='expense_change_other', user=self.li).exists())

    def test_create_forces_owner_to_requester(self):
        self.client.force_login(self.staff_a)
        response = self.client.post('/api/accounting/expenses/', {
            'expense_date': '2026-07-10', 'category': '交通費', 'amount': '800',
            'owner': self.staff_b.id, 'created_by': self.staff_b.id,
        }, content_type='application/json')
        self.assertEqual(response.status_code, 201, response.content)
        created = Expense.objects.get(pk=response.json()['id'])
        self.assertEqual(created.owner, self.staff_a)
        self.assertEqual(created.created_by, self.staff_a)

    # --- 集計・出力 -----------------------------------------------------------
    def test_summary_hides_balance_without_full_accounting(self):
        data = self.get(self.staff_a, '/api/accounting/expenses/summary/').json()
        self.assertEqual(data['total_expense'], 1000)
        self.assertIsNone(data['balance'])
        self.assertIsNone(data['opening_balance'])
        self.assertFalse(data['balance_visible'])
        self.assertEqual(data['expense_scope'], 'own')
        data = self.get(self.jiao, '/api/accounting/expenses/summary/').json()
        self.assertEqual(data['total_expense'], 10000)
        self.assertIsNone(data['balance'])
        self.assertEqual(data['expense_scope'], 'all')

    def test_summary_shows_balance_for_li(self):
        data = self.get(self.li, '/api/accounting/expenses/summary/').json()
        self.assertTrue(data['balance_visible'])
        self.assertEqual(data['balance'], 10000 - 10000)

    def test_excel_limited_to_own_without_export_all(self):
        from io import BytesIO

        from openpyxl import load_workbook

        for user, expected_total in ((self.staff_a, 1000), (self.jiao, 0)):
            response = self.get(user, '/api/accounting/expenses/excel/')
            self.assertEqual(response.status_code, 200)
            ws = load_workbook(BytesIO(response.content)).active
            values = [cell.value for row in ws.iter_rows() for cell in row if cell.value is not None]
            self.assertNotIn('期首残高', values)
            self.assertNotIn('B交通費', values)
            self.assertNotIn('李交通費', values)
            self.assertIn(expected_total, values)
        self.assertTrue(AuditLog.objects.filter(action='expense_export', user=self.jiao).exists())

    def test_excel_all_for_li(self):
        from io import BytesIO

        from openpyxl import load_workbook

        response = self.get(self.li, '/api/accounting/expenses/excel/')
        ws = load_workbook(BytesIO(response.content)).active
        values = [cell.value for row in ws.iter_rows() for cell in row if cell.value is not None]
        self.assertIn('期首残高', values)
        self.assertIn('B交通費', values)
        log = AuditLog.objects.get(action='expense_export', user=self.li)
        self.assertTrue(log.extra['includes_others'])

    def test_dashboard_scoped(self):
        data = self.get(self.staff_a, '/api/accounting/dashboard/').json()
        self.assertEqual(data['total_expense_amount'], 1000)
        self.assertIsNone(data['current_balance'])
        self.assertIsNone(data['total_income_source_amount'])
        self.assertEqual(data['recent_income_sources'], [])
        self.assertEqual({row['name'] for row in data['expense_category_chart']}, {'A交通費'})
        data = self.get(self.jiao, '/api/accounting/dashboard/').json()
        self.assertEqual(data['total_expense_amount'], 10000)
        self.assertIsNone(data['current_balance'])
        data = self.get(self.li, '/api/accounting/dashboard/').json()
        self.assertEqual(data['current_balance'], 0)

    def test_copy_expenses_rejects_unowned_reference(self):
        project = AccountingProject.objects.create(name='P')
        self.client.force_login(self.li)
        ok = self.client.post(f'/api/accounting/projects/{project.id}/copy-expenses/',
                              {'expense_ids': [self.exp_a.id]}, content_type='application/json')
        self.assertEqual(ok.status_code, 200)
        # 業務管理者は project モジュール自体が使えない
        self.client.force_login(self.jiao)
        denied = self.client.post(f'/api/accounting/projects/{project.id}/copy-expenses/',
                                  {'expense_ids': [self.exp_a.id]}, content_type='application/json')
        self.assertEqual(denied.status_code, 403)

    def test_copy_expenses_partial_scope_rejected(self):
        from django.contrib.auth.models import Permission

        project = AccountingProject.objects.create(name='P2')
        self.staff_a.user_permissions.add(Permission.objects.get(codename='use_project'))
        self.client.force_login(self.staff_a)
        response = self.client.post(f'/api/accounting/projects/{project.id}/copy-expenses/',
                                    {'expense_ids': [self.exp_a.id, self.exp_b.id]}, content_type='application/json')
        self.assertEqual(response.status_code, 400)
        self.assertEqual(project.project_expenses.count(), 0)

    def test_other_accounting_modules_only_for_li(self):
        for url in ('/api/accounting/income-sources/', '/api/accounting/vehicle-usages/',
                    '/api/accounting/projects/', '/api/accounting/vouchers/',
                    '/api/accounting/visa-return-applications/', '/api/accounting/tax-renewal-records/',
                    '/api/accounting/seifu-notice-records/'):
            self.assertEqual(self.get(self.jiao, url).status_code, 403, url)
            self.assertEqual(self.get(self.staff_a, url).status_code, 403, url)
            self.assertEqual(self.get(self.li, url).status_code, 200, url)

    def test_expense_category_read_for_expense_users_write_for_admin(self):
        self.assertEqual(self.get(self.staff_a, '/api/accounting/expense-categories/').status_code, 200)
        self.client.force_login(self.staff_a)
        response = self.client.post('/api/accounting/expense-categories/', {'name': '新分類'}, content_type='application/json')
        self.assertEqual(response.status_code, 403)
